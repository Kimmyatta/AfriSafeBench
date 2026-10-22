import os
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BLOG_DIR = ROOT / "docs" / "blog"


def configured_path(variable: str, default: Path) -> Path:
    value = os.environ.get(variable)
    if not value:
        return default
    path = Path(value)
    return path if path.is_absolute() else ROOT / path


SOURCE = configured_path(
    "AFRISAFEBENCH_BLOG_SOURCE",
    BLOG_DIR / "AfriSafeBench_Apart_Blog_Draft.md",
)
OUTPUT = configured_path(
    "AFRISAFEBENCH_BLOG_OUTPUT",
    BLOG_DIR / "AfriSafeBench_Apart_Blog_Submission.docx",
)
FOOTER_TEXT = os.environ.get(
    "AFRISAFEBENCH_BLOG_FOOTER",
    "AfriSafeBench | Apart Lab Studio Blog Submission",
)

INK = "17212B"
TEAL = "1F766D"
LIGHT_TEAL = "E8F3F1"
MUTED = "607086"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([color, underline])
    run.append(run_properties)

    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
)


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])

        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor.from_string(MUTED)
        else:
            link_match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link_match:
                add_hyperlink(paragraph, link_match.group(1), link_match.group(2))
        cursor = match.end()

    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(TEAL)
        style.font.bold = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(FOOTER_TEXT)
    footer_run.font.name = "Aptos"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(MUTED)

    document.core_properties.title = (
        "Evaluating LLM Recognition of AI Safety and Governance Risks in African Healthcare AI Deployments"
    )
    document.core_properties.author = "Kimberly Atta-Peters"
    document.core_properties.subject = "AfriSafeBench Apart Lab Studio blog post draft"


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    return rows


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, value)
            if row_index == 0:
                set_cell_shading(cell, LIGHT_TEAL)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(INK)
    document.add_paragraph()


def build_document() -> Path:
    document = Document()
    configure_document(document)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()

    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if not stripped or stripped == "---":
            index += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_table(document, parse_table(table_lines))
            continue

        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            image_path = (SOURCE.parent / image_match.group(2)).resolve()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Inches(6.2))
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption_run = caption.add_run(image_match.group(1))
            caption_run.italic = True
            caption_run.font.size = Pt(8.5)
            caption_run.font.color.rgb = RGBColor.from_string(MUTED)
            index += 1
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            title = heading_match.group(2)
            if level == 1:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, title)
            else:
                paragraph = document.add_heading(level=min(level - 1, 2))
                add_inline(paragraph, title)
            index += 1
            continue

        bullet_match = re.match(r"^-\s+(.+)$", stripped)
        if bullet_match:
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, bullet_match.group(1))
            index += 1
            continue

        numbered_match = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered_match:
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, numbered_match.group(1))
            index += 1
            continue

        paragraph = document.add_paragraph()
        add_inline(paragraph, stripped)
        if stripped.startswith("*Introducing AfriSafeBench"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith("**Kimberly Atta-Peters**"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        index += 1

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_document()
    print(output)
